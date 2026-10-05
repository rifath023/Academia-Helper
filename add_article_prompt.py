import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

PATH = r"D:\agent\academia blog writing\academia-helper-100-amazon-student-topics-2026-with-prompts.xlsx"
OUT = r"D:\agent\academia blog writing\academia-helper-100-amazon-student-topics-2026-with-prompts-v2.xlsx"
wb = openpyxl.load_workbook(PATH)
ws = wb["100 Amazon Topics"]
headers = [ws.cell(1,c).value for c in range(1, ws.max_column+1)]
idx_title = next(i+1 for i,h in enumerate(headers) if h and h.startswith("Article Title"))
idx_seed = headers.index("Seed Keywords") + 1
idx_words = headers.index("Suggested Words") + 1

new_col = ws.max_column + 1
hc = ws.cell(row=1, column=new_col, value="article writing prompt")
hc.fill = PatternFill("solid", fgColor="1C1917")
hc.font = Font(color="FFFFFF", bold=True, size=10)
hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(style="thin", color="D6D3D1")
hc.border = Border(left=thin, right=thin, top=thin, bottom=thin)

for r in range(2, ws.max_row+1):
    title = str(ws.cell(r, idx_title).value or "").strip()
    seed = str(ws.cell(r, idx_seed).value or "").strip()
    words = ws.cell(r, idx_words).value
    try:
        wc = str(int(words))
    except Exception:
        wc = str(words)
    prompt = (
        f"Suppose you are an expert SEO article writer who creates high-quality, natural, and Google-friendly content. "
        f"Write a {wc}-word article on '{title}' (seed keyword: {seed}).\n\n"
        f"Use proper grammar, clear language, natural keyword placement, and a reader-friendly structure. "
        f"Make the article informative, engaging, original, and easy for Google to understand and index.\n\n"
        f"Avoid unnecessary repetition, keyword stuffing, overly complex language, and filler content. "
        f"Keep the writing focused on the topic and make sure it provides genuine value to readers.\n\n"
        f"Article research data (paste the researched data from Amazon Research Prompt column here):\n"
        f"[PASTE RESEARCHED DATA HERE]"
    )
    cell = ws.cell(row=r, column=new_col, value=prompt)
    cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

from openpyxl.utils import get_column_letter
ws.column_dimensions[get_column_letter(new_col)].width = 90
for r in range(2, ws.max_row+1):
    if (ws.row_dimensions[r].height or 0) < 90:
        ws.row_dimensions[r].height = 90
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(new_col)}101"
wb.save(OUT)
print(f"saved col {new_col} rows {ws.max_row} -> {OUT}")
