import openpyxl, re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

PATH = r"D:\agent\academia blog writing\academia-helper-100-amazon-student-topics-2026.xlsx"
OUT = r"D:\agent\academia blog writing\academia-helper-100-amazon-student-topics-2026-with-prompts.xlsx"
wb = openpyxl.load_workbook(PATH)
ws = wb["100 Amazon Topics"]

# Find column indices by header
headers = [ws.cell(1,c).value for c in range(1, ws.max_column+1)]
print("Current headers:", headers)
idx_seed = headers.index("Seed Keywords") + 1
idx_title = next(i+1 for i,h in enumerate(headers) if h and h.startswith("Article Title"))
idx_slug = headers.index("URL Slug") + 1
idx_p1 = next(i+1 for i,h in enumerate(headers) if h and h.startswith("Product 1"))
idx_p5 = next(i+1 for i,h in enumerate(headers) if h and h.startswith("Product 5"))
idx_cat = headers.index("Category") + 1

def infer_budget(title, seed, cat):
    t = (title + " " + seed).lower()
    m = re.search(r"under\s*\$?\s*(\d+)", t)
    if m:
        return m.group(1)
    if "chair" in t or "standing desk" in t or "drafting table" in t or "printer" in t or "monitor" in t or "projector" in t:
        return "200"
    if "headphone" in t or "ssd" in t or "hard drive" in t or "tablet" in t or "webcam" in t or "microphone" in t:
        return "150"
    if "pen" in t or "highlight" in t or "sticky" in t or "flashcard" in t or "paper" in t or "binder" in t or "pencil" in t:
        return "30"
    if "book" in t or "pillow" in t or "glasses" in t or "shelf" in t or "organizer" in t or "mat" in t or "stand" in t:
        return "100"
    return "100"

new_col = ws.max_column + 1
header_cell = ws.cell(row=1, column=new_col, value="Amazon Research Prompt (copy-paste)")
header_cell.fill = PatternFill("solid", fgColor="1C1917")
header_cell.font = Font(color="FFFFFF", bold=True, size=10)
header_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(style="thin", color="D6D3D1")
header_cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

for r in range(2, ws.max_row+1):
    seed = str(ws.cell(r, idx_seed).value or "").strip()
    title = str(ws.cell(r, idx_title).value or "").strip()
    slug = str(ws.cell(r, idx_slug).value or "").strip()
    cat = str(ws.cell(r, idx_cat).value or "").strip()
    prods = [str(ws.cell(r, c).value or "").strip() for c in range(idx_p1, idx_p5+1)]
    prod_list = ", ".join([p for p in prods[:3] if p])
    budget = infer_budget(title, seed, cat)
    number = 5
    prompt = (
        f"Open Amazon and search '{seed}'. Based on this, find best {number} products "
        f"(reference examples: {prod_list}) which may attract buyers, budget within {budget} dollars.\n"
        f"Go to each product link and find out the pros and cons - it is better to get this information from customer reviews (3-star and 4-star reviews especially).\n"
        f"Collect each product's short description (2-3 lines, key specs: size, weight, material, compatibility).\n"
        f"Collect 1 image link per product (image URL only, not the image itself - preferably white-background main image).\n"
        f"Article title to use: '{title}' (slug: {slug}).\n"
        f"Search and draft 3-4 FAQs for the article (real buyer questions from Amazon Q&A / reviews about {seed})."
    )
    cell = ws.cell(row=r, column=new_col, value=prompt)
    cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

from openpyxl.utils import get_column_letter
ws.column_dimensions[get_column_letter(new_col)].width = 90
for r in range(2, ws.max_row+1):
    if ws.row_dimensions[r].height and ws.row_dimensions[r].height < 90:
        ws.row_dimensions[r].height = 90
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(new_col)}101"
wb.save(OUT)
print(f"saved, new col={new_col}, rows={ws.max_row} -> {OUT}")
