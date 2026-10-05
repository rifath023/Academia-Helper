import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

PATH = r"D:\agent\academia blog writing\academia-helper-100-amazon-student-topics-2026-with-prompts-v2.xlsx"
wb = openpyxl.load_workbook(PATH)
ws = wb["100 Amazon Topics"]
headers = [ws.cell(1,c).value for c in range(1, ws.max_column+1)]
idx_title = next(i+1 for i,h in enumerate(headers) if h and h.startswith("Article Title"))
idx_slug = headers.index("URL Slug") + 1

new_col = ws.max_column + 1
hc = ws.cell(row=1, column=new_col, value="html transform prompt")
hc.fill = PatternFill("solid", fgColor="1C1917")
hc.font = Font(color="FFFFFF", bold=True, size=10)
hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(style="thin", color="D6D3D1")
hc.border = Border(left=thin, right=thin, top=thin, bottom=thin)

for r in range(2, ws.max_row+1):
    title = str(ws.cell(r, idx_title).value or "").strip()
    slug = str(ws.cell(r, idx_slug).value or "").strip()
    prompt = (
        f"I have written the article for '{title}' (slug: {slug}).\n"
        f"[I will add the article here]\n\n"
        f"So, I need to transform this into HTML file for uploading into the site. "
        f"I have provided you a transactional demo format (file: 'transactional post format.html'), follow the format, and make ready and save in the path 'D:\\agent\\academia blog writing'. "
        f"Remember don't change a single word from the article, just exactly transform into the HTML format, and remember don't change the html format. Need exactly this format.\n\n"
        f"Here are all the related images\n"
        f"[here I will add the images]"
    )
    cell = ws.cell(row=r, column=new_col, value=prompt)
    cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

from openpyxl.utils import get_column_letter
ws.column_dimensions[get_column_letter(new_col)].width = 90
for r in range(2, ws.max_row+1):
    if (ws.row_dimensions[r].height or 0) < 90:
        ws.row_dimensions[r].height = 90
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(new_col)}101"
try:
    wb.save(PATH)
    print(f"saved in place col {new_col}")
except PermissionError:
    OUT = PATH.replace("-v2.xlsx", "-v3.xlsx")
    wb.save(OUT)
    print(f"locked, saved -> {OUT}")
